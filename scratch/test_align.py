import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.enum.text import WD_ALIGN_PARAGRAPH
import win32com.client, fitz, os, sys

sys.stdout.reconfigure(encoding='utf-8')

from pythainlp import word_tokenize
import re

NO_BREAK_BEFORE = r'[,.:;!?)\\]}\\%"\'\sฯๆ/”’]'
NO_BREAK_AFTER = r'[(\[{\s"\'/“‘]'

def thai_zwsp(text):
    if not text:
        return ""
    lines = str(text).split('\n')
    processed = []
    for line in lines:
        if not line:
            processed.append("")
            continue
        tokens = word_tokenize(line, engine='newmm')
        joined = '\u200b'.join(tokens)
        joined = re.sub(r'\u200b+(' + NO_BREAK_BEFORE + ')', r'\1', joined)
        joined = re.sub('(' + NO_BREAK_AFTER + r')\u200b+', r'\1', joined)
        joined = re.sub(r'\u200b+', '\u200b', joined)
        processed.append(joined)
    return '\n'.join(processed)

doc = docx.Document()
s = doc.sections[0]
s.left_margin = docx.shared.Inches(1.5)
s.right_margin = docx.shared.Inches(1.0)

text_intro = "โครงสร้างและการจัดวางสูตรคำนวณอัตโนมัติในไฟล์สมุดบันทึกสุขภาพ '6806021612037_Health30D.xlsx' บน Microsoft OneDrive ประกอบด้วย:"

# Test 1: thaiDistribute (current)
p1_hdr = doc.add_paragraph("--- 1. thaiDistribute ---")
p1 = doc.add_paragraph()
p1.alignment = WD_ALIGN_PARAGRAPH.THAI_JUSTIFY
p1.add_run(thai_zwsp(text_intro))

# Test 2: both (Standard Justify)
p2_hdr = doc.add_paragraph("--- 2. both (Justify) ---")
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p2.add_run(thai_zwsp(text_intro))

# Test 3: left
p3_hdr = doc.add_paragraph("--- 3. left ---")
p3 = doc.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.LEFT
p3.add_run(thai_zwsp(text_intro))

# Test 4: Biography cell simulation
doc.add_page_break()
doc.add_paragraph("--- Biography cell test ---")
tbl = doc.add_table(rows=1, cols=2)
c0 = tbl.cell(0, 0)
c1 = tbl.cell(0, 1)
c0.width = docx.shared.Inches(1.5)
c1.width = docx.shared.Inches(4.25)
c0.paragraphs[0].text = "สถานที่ติดต่อ:"

# Try c1 with left vs thaiDistribute
p_c1 = c1.paragraphs[0]
p_c1.alignment = WD_ALIGN_PARAGRAPH.LEFT
p_c1.add_run(thai_zwsp("คณะเทคโนโลยีและการจัดการอุตสาหกรรม มหาวิทยาลัยเทคโนโลยีพระจอมเกล้าพระนครเหนือ วิทยาเขตปราจีนบุรี 129 หมู่ 21 ตำบลเนินหอม อำเภอเมือง จังหวัดปราจีนบุรี 25230"))

doc.save('scratch/test_align.docx')
word = win32com.client.Dispatch('Word.Application')
wdoc = word.Documents.Open(os.path.abspath('scratch/test_align.docx'))
wdoc.SaveAs(os.path.abspath('scratch/test_align.pdf'), FileFormat=17)
wdoc.Close(False)
word.Quit()

fdoc = fitz.open('scratch/test_align.pdf')
pix1 = fdoc[0].get_pixmap(dpi=150)
pix1.save('scratch/test_align_p1.png')
pix2 = fdoc[1].get_pixmap(dpi=150)
pix2.save('scratch/test_align_p2.png')
print("Rendered test pages to scratch/test_align_p1.png and p2.png")
