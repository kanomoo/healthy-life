import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import win32com.client, fitz, os, sys

sys.stdout.reconfigure(encoding='utf-8')

from pythainlp import word_tokenize
import re

NO_BREAK_BEFORE = r'[,.:;!?)\\]}\\%"\'\sฯๆ/”’]'
NO_BREAK_AFTER = r'[(\[{\s"\'/“‘]'

def thai_zwsp(text):
    if not text:
        return ""
    lines = str(text).split('\n')
    processed = []
    for line in lines:
        if not line:
            processed.append("")
            continue
        tokens = word_tokenize(line, engine='newmm')
        joined = '\u200b'.join(tokens)
        joined = re.sub(r'\u200b+(' + NO_BREAK_BEFORE + ')', r'\1', joined)
        joined = re.sub('(' + NO_BREAK_AFTER + r')\u200b+', r'\1', joined)
        joined = re.sub(r'\u200b+', '\u200b', joined)
        processed.append(joined)
    return '\n'.join(processed)

def set_run_font(run, size_pt=14, bold=False):
    run.font.name = "TH Sarabun PSK"
    run.font.size = docx.shared.Pt(size_pt)
    run.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(docx.oxml.ns.qn('w:ascii'), "TH Sarabun PSK")
    rFonts.set(docx.oxml.ns.qn('w:hAnsi'), "TH Sarabun PSK")
    rFonts.set(docx.oxml.ns.qn('w:cs'), "TH Sarabun PSK")
    sz_val = int(round(size_pt * 2))
    rPr.append(parse_xml(f'<w:sz {nsdecls("w")} w:val="{sz_val}"/>'))
    rPr.append(parse_xml(f'<w:szCs {nsdecls("w")} w:val="{sz_val}"/>'))
    if bold:
        rPr.append(parse_xml(f'<w:b {nsdecls("w")}/>'))
        rPr.append(parse_xml(f'<w:bCs {nsdecls("w")}/>'))

doc = docx.Document()
s = doc.sections[0]
s.left_margin = docx.shared.Inches(1.5)
s.right_margin = docx.shared.Inches(1.0)
s.top_margin = docx.shared.Inches(1.5)
s.bottom_margin = docx.shared.Inches(1.0)

# Appendix D Header
p_t = doc.add_paragraph()
p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p_t.add_run(thai_zwsp("ภาคผนวก ง\nโครงสร้างไฟล์ Microsoft Excel และสูตรคำนวณที่ใช้ในการวิจัย\n(Excel Workbook Architecture & Formulas)"))
set_run_font(r, size_pt=16, bold=True)

# Intro paragraph with LEFT alignment
p_intro = doc.add_paragraph()
p_intro.alignment = WD_ALIGN_PARAGRAPH.LEFT
p_intro.paragraph_format.space_before = docx.shared.Pt(6)
p_intro.paragraph_format.space_after = docx.shared.Pt(4)
r_in = p_intro.add_run(thai_zwsp("โครงสร้างและการจัดวางสูตรคำนวณอัตโนมัติในไฟล์สมุดบันทึกสุขภาพ '6806021612037_Health30D.xlsx' บนระบบคลาวด์ Microsoft OneDrive ประกอบด้วย 3 เวิร์กชีตหลัก และสูตรคำนวณมาตรฐานดังแสดงในตารางที่ ง.1:"))
set_run_font(r_in, size_pt=14)

# Table caption
p_cap = doc.add_paragraph()
p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
p_cap.paragraph_format.space_before = docx.shared.Pt(6)
p_cap.paragraph_format.space_after = docx.shared.Pt(2)
r_c1 = p_cap.add_run("ตารางที่ ง.1  ")
set_run_font(r_c1, size_pt=13, bold=True)
r_c2 = p_cap.add_run(thai_zwsp("รายละเอียดโครงสร้างเวิร์กชีตและสูตรคำนวณอัตโนมัติในไฟล์ Microsoft Excel"))
set_run_font(r_c2, size_pt=13, bold=False)

# Table
headers = ["ตัวชี้วัด / รายการ", "เวิร์กชีต", "สูตรคำนวณ Excel (Formula Syntax)", "คำอธิบายการประมวลผล"]
data = [
    ["รอบพักที่ควรทำ", "Daily_Logbook", "=INT(E2/45)", "หารเวลาใช้จอ (นาที) ด้วย 45 ปัดเศษทิ้ง"],
    ["ร้อยละการพักจริง", "Daily_Logbook", '=IF(F2>0, (G2/F2)*100, "-")', "สัดส่วนรอบพักจริงต่อรอบที่ควรทำ"],
    ["ลดปวดทันทีหลังนวด", "Daily_Logbook", "=IF(I2>0, ((I2-J2)/I2)*100, 0)", "อัตราลดลงของ NRS หลังนวดเทียบก่อนนวด"],
    ["ลดปวดสุทธิ (KR 1)", "Summary_Metrics", "=((I2-AVERAGE(I13:I15))/I2)*100", "เปรียบเทียบ Baseline D1 กับเฉลี่ย D12–D14"],
    ["ความสม่ำเสมอ (KR 4)", "Summary_Metrics", '=(COUNTIF(L2:L15, "ครบ")/14)*100', "สัดส่วนวันที่ทำโปรแกรมนวดและบริหารครบ"],
    ["มุมข้อต่อสรีระ 3 จุด", "Angle_Analysis", "ImageMeter Link / Input Cells", "บันทึกค่ามุมศอก สะโพก เข่า จากภาพถ่าย"]
]
col_w = [docx.shared.Inches(1.25), docx.shared.Inches(1.05), docx.shared.Inches(1.85), docx.shared.Inches(1.6)]

tbl = doc.add_table(rows=len(data)+1, cols=4)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl.autofit = False

tblPr = tbl._tbl.tblPr
tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="8280" w:type="dxa"/>'))
tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
tblBorders = parse_xml(
    f'<w:tblBorders {nsdecls("w")}>'
    f'<w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
    f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
    f'<w:left w:val="none"/>'
    f'<w:right w:val="none"/>'
    f'<w:insideH w:val="none"/>'
    f'<w:insideV w:val="none"/>'
    f'</w:tblBorders>'
)
tblPr.append(tblBorders)

# Header row
for c_idx, h_text in enumerate(headers):
    cell = tbl.cell(0, c_idx)
    cell.width = col_w[c_idx]
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = docx.shared.Pt(2)
    p.paragraph_format.space_after = docx.shared.Pt(2)
    r = p.add_run(thai_zwsp(h_text))
    set_run_font(r, size_pt=11, bold=True)
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/></w:tcBorders>')
    tcPr.append(tcBorders)

for r_idx, row in enumerate(data):
    for c_idx, val in enumerate(row):
        cell = tbl.cell(r_idx+1, c_idx)
        cell.width = col_w[c_idx]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx != 1 else WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = docx.shared.Pt(2)
        p.paragraph_format.space_after = docx.shared.Pt(2)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(thai_zwsp(val))
        set_run_font(r, size_pt=10, bold=False)

# Add Biography Page
doc.add_page_break()
p_bio_t = doc.add_paragraph()
p_bio_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_bt = p_bio_t.add_run(thai_zwsp("ประวัติผู้จัดทำ"))
set_run_font(r_bt, size_pt=18, bold=True)

bio_data = [
    ["ชื่อ - นามสกุล:", "นายปภาวิน ธิติชุณหกุล (Mr. Paphavin Thitichunhakun)"],
    ["รหัสนักศึกษา:", "6806021612037"],
    ["การศึกษา:", "นักศึกษาชั้นปีที่ 1 แขนงวิชาเทคโนโลยีสารสนเทศ (IT)\nภาควิชาเทคโนโลยีสารสนเทศ คณะเทคโนโลยีและการจัดการอุตสาหกรรม\nมหาวิทยาลัยเทคโนโลยีพระจอมเกล้าพระนครเหนือ วิทยาเขตปราจีนบุรี"],
    ["รายวิชา:", "080303609 สุขภาพเพื่อชีวิต (Healthy Life) กลุ่มเรียนที่ 2 (Sec 2)\nภาคการศึกษาที่ 1 ปีการศึกษา 2569"],
    ["สถานที่ติดต่อ:", "คณะเทคโนโลยีและการจัดการอุตสาหกรรม มหาวิทยาลัยเทคโนโลยีพระจอมเกล้าพระนครเหนือ\nวิทยาเขตปราจีนบุรี 129 หมู่ 21 ตำบลเนินหอม อำเภอเมือง จังหวัดปราจีนบุรี 25230"],
    ["ความสนใจทางวิชาการ:", "การยศาสตร์เชิงคอมพิวเตอร์ (Computer Ergonomics), การวิเคราะห์ข้อมูลสุขภาพเชิงประจักษ์ (Empirical Health Analytics), และการจัดการสารสนเทศส่วนบุคคล (Personal Information Management)"]
]

tbl_bio = doc.add_table(rows=len(bio_data), cols=2)
tbl_bio.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl_bio.autofit = False
tbl_bioPr = tbl_bio._tbl.tblPr
tbl_bioPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="8280" w:type="dxa"/>'))
tbl_bioPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
tbl_bioBorders = parse_xml(
    f'<w:tblBorders {nsdecls("w")}>'
    f'<w:top w:val="none"/><w:bottom w:val="none"/><w:left w:val="none"/>'
    f'<w:right w:val="none"/><w:insideH w:val="none"/><w:insideV w:val="none"/>'
    f'</w:tblBorders>'
)
tbl_bioPr.append(tbl_bioBorders)

col_w_bio = [docx.shared.Inches(1.5), docx.shared.Inches(4.25)]

for r_idx, (lbl, val) in enumerate(bio_data):
    cell_lbl = tbl_bio.cell(r_idx, 0)
    cell_val = tbl_bio.cell(r_idx, 1)
    cell_lbl.width = col_w_bio[0]
    cell_val.width = col_w_bio[1]
    
    p0 = cell_lbl.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p0.paragraph_format.space_before = docx.shared.Pt(2)
    p0.paragraph_format.space_after = docx.shared.Pt(3)
    r0 = p0.add_run(thai_zwsp(lbl))
    set_run_font(r0, size_pt=15, bold=True)
    
    lines = val.split('\n')
    for l_idx, line in enumerate(lines):
        if l_idx == 0:
            p1 = cell_val.paragraphs[0]
        else:
            p1 = cell_val.add_paragraph()
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p1.paragraph_format.space_before = docx.shared.Pt(0)
        p1.paragraph_format.space_after = docx.shared.Pt(2)
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(thai_zwsp(line))
        set_run_font(r1, size_pt=15, bold=False)

doc.save('scratch/test_final_2pages.docx')
word = win32com.client.Dispatch('Word.Application')
wdoc = word.Documents.Open(os.path.abspath('scratch/test_final_2pages.docx'))
wdoc.SaveAs(os.path.abspath('scratch/test_final_2pages.pdf'), FileFormat=17)
wdoc.Close(False)
word.Quit()

fdoc = fitz.open('scratch/test_final_2pages.pdf')
fdoc[0].get_pixmap(dpi=150).save('scratch/test_final_p1.png')
fdoc[1].get_pixmap(dpi=150).save('scratch/test_final_p2.png')
print("Rendered scratch/test_final_p1.png and p2.png")
